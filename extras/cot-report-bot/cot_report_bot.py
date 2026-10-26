#!/usr/bin/env python3
"""
CFTC COT Disaggregated Report Bot
Fetches weekly petroleum COT data, formats a summary table and AI commentary,
and sends both to Telegram.
"""

import os
import re
import json
import base64
import requests
import anthropic
from datetime import datetime, timedelta, timezone
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────
TELEGRAM_BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
TELEGRAM_CHAT_ID   = os.environ["TELEGRAM_CHAT_ID"]
ANTHROPIC_API_KEY  = os.environ["ANTHROPIC_API_KEY"]

CFTC_URL = "https://www.cftc.gov/dea/options/petroleum_lof.htm"

# Model for both the news search and the commentary. Override with COT_BOT_MODEL.
MODEL = os.getenv("COT_BOT_MODEL") or "claude-sonnet-5-5"

# ── Saving the commentary (read by the Streamlit dashboard) ───────────────────
# One JSON file per report date, e.g. commentary/2026-09-29.json.
COMMENTARY_DIR = Path(os.getenv("COMMENTARY_DIR", "commentary"))
# This page is futures + options combined (the Excel history below filters on
# "Combined" too). If you ever point the bot at the futures-only page, change this.
COT_BASIS = "combined"
# Optional: also push each file to a GitHub repo so the dashboard can read it.
# Needs a fine-grained token with Contents: read & write on that one repo.
GITHUB_TOKEN  = os.getenv("GITHUB_TOKEN")           # leave unset to save locally only
GITHUB_REPO   = os.getenv("COMMENTARY_REPO")        # "owner/repo"
GITHUB_BRANCH = os.getenv("COMMENTARY_BRANCH", "main")
GITHUB_FOLDER = os.getenv("COMMENTARY_FOLDER", "commentary")
# Set to 1 on a schedule: if this week's commentary is already saved, do nothing.
# Stops a holiday-delayed CFTC release from re-sending last week's messages.
SKIP_IF_ALREADY_SAVED = os.getenv("SKIP_IF_ALREADY_SAVED") == "1"

TARGET_CONTRACTS = [
    "BRENT LAST DAY",
    "WTI FINANCIAL CRUDE OIL",
    "GASOLINE RBOB",
    "NY HARBOR ULSD",
    "GULF JET NY HEAT OIL SPR",
    "CRUDE OIL, LIGHT SWEET",   # WTI ICE
    "WTI-PHYSICAL",
]

SHORT_NAMES = {
    "BRENT LAST DAY":           "Brent LD (NYMEX)",
    "WTI FINANCIAL CRUDE OIL":  "WTI Fin (NYMEX)",
    "GASOLINE RBOB":            "RBOB (NYMEX)",
    "NY HARBOR ULSD":           "ULSD (NYMEX)",
    "GULF JET NY HEAT OIL SPR": "Jet/HO Spr (NYMEX)",
    "CRUDE OIL, LIGHT SWEET":   "WTI (ICE)",
    "WTI-PHYSICAL":             "WTI Phys (NYMEX)",
}

CAT_ROWS = [
    ("Prod/Merch", "pm_long", "pm_short", "d_pm_long", "d_pm_short"),
    ("Swap Dlrs",  "sd_long", "sd_short", "d_sd_long", "d_sd_short"),
    ("Mgd Money",  "mm_long", "mm_short", "d_mm_long", "d_mm_short"),
    ("Other Rpt",  "or_long", "or_short", "d_or_long", "d_or_short"),
]

# ── Fetch ─────────────────────────────────────────────────────────────────────

def fetch_html(url):
    headers = {"User-Agent": "Mozilla/5.0 (COT-Bot/1.0)"}
    r = requests.get(url, headers=headers, timeout=30)
    r.raise_for_status()
    return r.text

# ── Parse ─────────────────────────────────────────────────────────────────────

def extract_numbers(line):
    return [int(n.replace(",", "")) for n in re.findall(r"-?\d[\d,]*", line)]

def parse_report(html):
    soup = BeautifulSoup(html, "html.parser")
    pre  = soup.find("pre")
    if not pre:
        raise ValueError("No <pre> block found in CFTC HTML")

    contracts   = {}
    cur         = None
    report_date = None
    in_pos      = False
    in_chg      = False
    pos_done    = False
    chg_done    = False

    for line in pre.get_text().splitlines():
        if "Code-" in line and "Disaggregated" not in line:
            cur = line.split("Code-")[0].strip()
            in_pos = in_chg = pos_done = chg_done = False
            continue

        if cur and "Disaggregated Commitments of Traders" in line:
            m = re.search(r",\s+(\w+ \d+, \d{4})", line)
            if m:
                report_date = m.group(1)
            continue

        if "Positions" in line and "Percent" not in line and "Number" not in line and "Changes" not in line:
            in_pos = True; in_chg = False
            continue

        if "Changes in Commitments" in line:
            in_chg = True; in_pos = False
            continue

        if "Percent of Open Interest" in line or "Number of Traders" in line:
            in_pos = in_chg = False
            continue

        if in_pos and not pos_done and cur and re.match(r"\s*All\s*:", line):
            nums = extract_numbers(line)
            if len(nums) >= 12:
                if cur not in contracts:
                    contracts[cur] = {}
                contracts[cur].update({
                    "oi": nums[0],
                    "pm_long": nums[1],  "pm_short": nums[2],
                    "sd_long": nums[3],  "sd_short": nums[4],
                    "mm_long": nums[6],  "mm_short": nums[7],
                    "or_long": nums[9],  "or_short": nums[10],
                })
                pos_done = True
            continue

        if in_chg and not chg_done and cur and cur in contracts:
            nums = extract_numbers(line)
            if len(nums) >= 12:
                contracts[cur].update({
                    "d_oi": nums[0],
                    "d_pm_long": nums[1],  "d_pm_short": nums[2],
                    "d_sd_long": nums[3],  "d_sd_short": nums[4],
                    "d_mm_long": nums[6],  "d_mm_short": nums[7],
                    "d_or_long": nums[9],  "d_or_short": nums[10],
                })
                chg_done = True
            continue

    return contracts, report_date

# ── Contract Matching ─────────────────────────────────────────────────────────

def match_contracts(parsed):
    matched = {}
    for target in TARGET_CONTRACTS:
        for key in parsed:
            if target.upper() in key.upper():
                matched[target] = parsed[key]
                break
    return matched

# ── Formatting helpers ────────────────────────────────────────────────────────

def f(n, sign=False):
    if n is None:
        return "N/A"
    s = f"{abs(n):,}"
    if sign:
        return ("+" if n >= 0 else "-") + s
    return ("-" if n < 0 else "") + s

# ── Message 1: Table ──────────────────────────────────────────────────────────

def build_table_message(matched, report_date):
    lines = []
    lines.append(f"📊 *CFTC COT Disaggregated — Petroleum*")
    lines.append(f"_Week of {report_date}_\n")

    for target, data in matched.items():
        cname = SHORT_NAMES.get(target, target[:18])
        oi    = data.get("oi", 0)
        d_oi  = data.get("d_oi", 0)

        lines.append(f"*── {cname} ──*")
        lines.append(f"_OI: {f(oi)} ({f(d_oi, sign=True)})_")
        lines.append("```")
        lines.append(f"{'Category':<11} {'Long':>8} {'Short':>8} {'Net':>8} {'ΔNet':>8}")
        lines.append("-" * 47)

        tot_long = tot_short = tot_dnt = 0

        for label, lk, sk, dlk, dsk in CAT_ROWS:
            lg  = data.get(lk, 0)
            sh  = data.get(sk, 0)
            nt  = lg - sh
            dnt = data.get(dlk, 0) - data.get(dsk, 0)
            lines.append(
                f"{label:<11} {f(lg):>8} {f(sh):>8} {f(nt):>8} {f(dnt, sign=True):>8}"
            )
            tot_long  += lg
            tot_short += sh
            tot_dnt   += dnt

        tot_net = tot_long - tot_short
        lines.append("-" * 47)
        lines.append(
            f"{'TOTAL':<11} {f(tot_long):>8} {f(tot_short):>8} {f(tot_net):>8} {f(tot_dnt, sign=True):>8}"
        )
        lines.append("```\n")

    return "\n".join(lines)

# ── News via Claude web search ───────────────────────────────────────────────

def fetch_oil_news():
    """
    Use the Anthropic API with web_search tool to fetch broad oil market news
    plus macro/fundamental drivers. Returns a news summary string.
    """
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    search_prompt = """Search for the latest news (past 7 days) on the following topics and return a concise summary of the key developments:

1. Crude oil market news — Brent and WTI price moves, any significant rallies or selloffs and their drivers
2. OPEC+ — any production decisions, quota changes, compliance issues, or member commentary
3. US EIA inventory data — latest weekly crude, gasoline, and distillate inventory changes (draws or builds)
4. Oil demand — China demand outlook, global refinery runs, any demand surprises
5. Geopolitical factors — anything affecting oil supply (Middle East tensions, Russia/Ukraine, sanctions, shipping disruptions)
6. Macro backdrop — US dollar direction, Fed rate expectations, and how they are influencing oil positioning

For each topic, give 2-3 bullet points of the most important developments. Be specific — include price levels, inventory figures in millions of barrels, and named officials or countries where relevant."""

    try:
        response = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            tools=[{"type": "web_search_20250305", "name": "web_search"}],
            messages=[{"role": "user", "content": search_prompt}]
        )
        news_text = text_of(response)
        return news_text if news_text else ""
    except Exception as e:
        print(f"News fetch error: {e}")
        return ""

# ── Claude Commentary ─────────────────────────────────────────────────────────

def text_of(response):
    """The written answer only. Newer models can put a thinking block first, which
    has no .text, so take the blocks of type "text" and ignore the rest."""
    return "".join(
        block.text for block in response.content if getattr(block, "type", "") == "text"
    ).strip()


def build_commentary(matched, report_date, news):
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    data_lines = []
    for target, d in matched.items():
        cname = SHORT_NAMES.get(target, target)
        oi    = d.get("oi", 0)
        data_lines.append(f"\n{cname} (OI: {oi:,}):")
        for label, lk, sk, dlk, dsk in CAT_ROWS:
            lg  = d.get(lk, 0); sh = d.get(sk, 0); nt = lg - sh
            dnt = d.get(dlk, 0) - d.get(dsk, 0)
            data_lines.append(
                f"  {label}: L={lg:,} S={sh:,} Net={nt:,} (ΔNet={'+' if dnt>=0 else ''}{dnt:,})"
            )

    news_section = f"\nRECENT OIL MARKET NEWS:\n{news}\n" if news.strip() else ""

    news_instruction = (
        "use the market and macro context above to explain the positioning moves — "
        "explicitly connect OPEC+ decisions, inventory data, demand signals, "
        "geopolitical developments, and dollar/rates direction to what speculators "
        "and commercials are doing across the petroleum complex"
        if news.strip() else
        "write a concise two-liner on the fundamental backdrop based on positioning "
        "signals alone. Do not flag caveats or suggest caution about missing data."
    )

    prompt = f"""You are a senior commodities analyst covering energy markets with deep knowledge of oil market fundamentals and macro drivers.

Analyse the CFTC Disaggregated COT data below for the week of {report_date} and write a
sharp, professional commentary (max 350 words) for a petroleum trading desk audience.

COT DATA:
{''.join(data_lines)}
{news_section}
Structure your commentary with these sections:
🔑 Key Positioning Shifts — biggest WoW moves across the petroleum complex, especially Managed Money. Use specific numbers. Note any divergences between crude benchmarks (Brent vs WTI) or between crude and products (RBOB, ULSD).
⚖️ Physical vs Financial — divergences between Prod/Merch hedgers and Managed Money/Swaps. Are commercials hedging into strength or weakness?
🌐 Macro & Fundamentals — {news_instruction}. Tie together OPEC+ supply policy, inventory trends, demand signals, and macro (dollar, rates) to explain the positioning. Be specific about cause and effect.
⚡ Risk Outlook — given positioning and fundamentals, state the near-term price risk for crude and products. Flag any crowded positioning risk or potential for a sharp positioning unwind.

Be direct and actionable. No waffle. Use specific numbers from the data and news.

End your commentary with the line: \"— End —\""""  

    msg = client.messages.create(
        model=MODEL,
        max_tokens=4000,
        messages=[{"role": "user", "content": prompt}]
    )
    text = text_of(msg)
    if not text:
        raise RuntimeError(f"The model returned no text (stop reason: {msg.stop_reason}).")
    return text

# ── Save commentary for the dashboard ─────────────────────────────────────────

def build_record(matched, report_date, news, commentary):
    """Everything the dashboard needs to show this week's commentary, and to check
    it against its own numbers."""
    this_week = datetime.strptime(report_date, "%B %d, %Y").date()
    contracts = {}
    for target, d in matched.items():
        cats = {}
        for label, lk, sk, dlk, dsk in CAT_ROWS:
            lg, sh = d.get(lk, 0), d.get(sk, 0)
            dl, ds = d.get(dlk, 0), d.get(dsk, 0)
            cats[label] = {"long": lg, "short": sh, "net": lg - sh,
                           "d_long": dl, "d_short": ds, "d_net": dl - ds}
        contracts[SHORT_NAMES.get(target, target)] = {
            "cftc_name": target, "oi": d.get("oi", 0), "d_oi": d.get("d_oi", 0),
            "categories": cats,
        }
    return {
        "schema": 1,
        "report": "disaggregated",
        "basis": COT_BASIS,
        "report_date": this_week.isoformat(),
        # The CFTC "changes" rows are versus the previous weekly report.
        "prior_date": (this_week - timedelta(days=7)).isoformat(),
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "model": MODEL,
        "commentary": commentary.replace("— End —", "").strip(),
        "news": news.strip(),
        "contracts": contracts,
    }


def already_saved(report_date):
    try:
        day = datetime.strptime(report_date, "%B %d, %Y").date().isoformat()
    except (TypeError, ValueError):
        return False
    return (COMMENTARY_DIR / f"{day}.json").exists()


def push_to_github(path_in_repo, content, message):
    api = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{path_in_repo}"
    headers = {"Authorization": f"Bearer {GITHUB_TOKEN}",
               "Accept": "application/vnd.github+json"}
    body = {"message": message, "branch": GITHUB_BRANCH,
            "content": base64.b64encode(content).decode()}
    existing = requests.get(api, headers=headers, params={"ref": GITHUB_BRANCH}, timeout=15)
    if existing.status_code == 200:            # re-run for the same week: overwrite
        body["sha"] = existing.json()["sha"]
    requests.put(api, headers=headers, json=body, timeout=20).raise_for_status()


def save_commentary(matched, report_date, news, commentary):
    """Never raises: a storage problem must not stop the Telegram messages."""
    try:
        if not report_date or not commentary.strip():
            print("Commentary not saved (no report date or empty commentary).")
            return
        record = build_record(matched, report_date, news, commentary)
        COMMENTARY_DIR.mkdir(parents=True, exist_ok=True)
        path = COMMENTARY_DIR / f"{record['report_date']}.json"
        payload = json.dumps(record, indent=2, ensure_ascii=False).encode("utf-8")
        path.write_bytes(payload)
        print(f"Commentary saved: {path}")
        if GITHUB_TOKEN and GITHUB_REPO:
            push_to_github(f"{GITHUB_FOLDER}/{path.name}", payload,
                           f"COT commentary {record['report_date']}")
            print(f"Commentary pushed to {GITHUB_REPO}.")
    except Exception as e:
        print(f"Commentary save error: {e}")

# ── Telegram ──────────────────────────────────────────────────────────────────

def split_message(text, limit=4000):
    """Split text into chunks at sentence/paragraph boundaries to avoid mid-sentence cuts."""
    if len(text) <= limit:
        return [text]
    chunks = []
    while len(text) > limit:
        # Try to split at paragraph break first
        split_at = text.rfind("\n\n", 0, limit)
        if split_at == -1:
            # Fall back to sentence end
            split_at = max(
                text.rfind(". ", 0, limit),
                text.rfind("! ", 0, limit),
                text.rfind("? ", 0, limit),
                text.rfind("\n", 0, limit),
            )
        if split_at == -1:
            # Last resort: hard cut
            split_at = limit
        else:
            split_at += 1  # include the punctuation
        chunks.append(text[:split_at].strip())
        text = text[split_at:].strip()
    if text:
        chunks.append(text)
    return chunks


def send_telegram(text, parse_mode="Markdown"):
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    for chunk in split_message(text):
        r = requests.post(url, json={
            "chat_id":    TELEGRAM_CHAT_ID,
            "text":       chunk,
            "parse_mode": parse_mode,
        }, timeout=15)
        if r.status_code != 200:
            print(f"Telegram error {r.status_code}: {r.text}")
            print(f"--- Failed chunk ---\n{chunk}\n--- End ---")
        r.raise_for_status()


# ── Excel generation ──────────────────────────────────────────────────────────

def fetch_historical_cot(match_str, weeks=13):
    """Fetch last 13 weeks of COT data from CFTC Socrata API."""
    date_from = (datetime.utcnow() - timedelta(days=weeks * 7 + 14)).strftime("%Y-%m-%d")
    all_rows  = []
    offset    = 0
    while True:
        params = {
            "$where":  (f"report_date_as_yyyy_mm_dd >= '{date_from}' "
                        f"AND upper(market_and_exchange_names) like '{match_str.upper()}%' "
                        f"AND futonly_or_combined = 'Combined'"),
            "$order":  "report_date_as_yyyy_mm_dd DESC",
            "$limit":  500,
            "$offset": offset,
        }
        try:
            r = requests.get("https://publicreporting.cftc.gov/resource/kh3c-gbw2.json",
                             params=params, timeout=30)
            r.raise_for_status()
            rows = r.json()
        except Exception as e:
            print(f"    API error for {match_str}: {e}")
            break
        if not rows:
            break
        all_rows.extend(rows)
        if len(rows) < 500:
            break
        offset += 500

    # Debug: print actual values from first row
    if all_rows:
        sample = all_rows[0]
        print(f"    Contract: {sample.get('market_and_exchange_names')}")
        print(f"    fut_only_or_combined: {sample.get('futonly_or_combined')}")
        for k, v in sorted(sample.items()):
            if any(x in k for x in ["prod","swap","money","other_rept","open_interest"]):
                print(f"      {k}: {v}")

    results = []
    for row in all_rows[:weeks]:
        try:
            # PM: no _all suffix (confirmed from API debug)
            pm_long  = int(float(row.get("prod_merc_positions_long",  row.get("prod_merc_positions_long_all",  0)) or 0))
            pm_short = int(float(row.get("prod_merc_positions_short", row.get("prod_merc_positions_short_all", 0)) or 0))
            # SD: long has _all, short has double underscore
            sd_long  = int(float(row.get("swap_positions_long_all",   0) or 0))
            sd_short = int(float(row.get("swap__positions_short_all", row.get("swap_positions_short_all", 0)) or 0))
            # MM: both have _all
            mm_long  = int(float(row.get("m_money_positions_long_all",  0) or 0))
            mm_short = int(float(row.get("m_money_positions_short_all", 0) or 0))
            # OR: no _all suffix (confirmed from API debug)
            or_long  = int(float(row.get("other_rept_positions_long",  row.get("other_rept_positions_long_all",  0)) or 0))
            or_short = int(float(row.get("other_rept_positions_short", row.get("other_rept_positions_short_all", 0)) or 0))
            oi       = int(float(row.get("open_interest_all", 0) or 0))

            results.append({
                "date":     row.get("report_date_as_yyyy_mm_dd", "")[:10],
                "pm_long":  pm_long,  "pm_short": pm_short,
                "sd_long":  sd_long,  "sd_short": sd_short,
                "mm_long":  mm_long,  "mm_short": mm_short,
                "or_long":  or_long,  "or_short": or_short,
                "oi":       oi,
            })
        except Exception as e:
            print(f"    Row parse error: {e}")
            continue
    return results


def build_excel(matched, report_date, bot_label):
    """Build single-sheet Excel with all contracts stacked vertically."""
    NAVY   = "1F3864"
    BLUE   = "2E75B6"
    WHITE  = "FFFFFF"
    ALTROW = "EEF4FB"
    GREY   = "F2F2F2"

    def xfill(hex_color): return PatternFill("solid", fgColor=hex_color)
    def xfont(bold=False, size=9, color="000000", italic=False):
        return Font(name="Arial", bold=bold, size=size, color=color, italic=italic)
    def xaln(h="center"): return Alignment(horizontal=h, vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    def xbdr(): return Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "COT Data"
    ws.freeze_panes = "A4"

    COLS = [
        ("Date", 14),
        ("PM Long", 9), ("PM Short", 9),
        ("SD Long", 9), ("SD Short", 9),
        ("MM Long", 9), ("MM Short", 9),
        ("OR Long", 9), ("OR Short", 9),
        ("OI", 12),
    ]
    NCOLS = len(COLS)
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Row 1: title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOLS)
    c = ws.cell(1, 1, f"CFTC COT Disaggregated — {bot_label} | Week of {report_date}")
    c.font = Font(name="Arial", bold=True, size=13, color=NAVY)
    c.alignment = xaln("left")
    ws.row_dimensions[1].height = 22

    # Row 2: group headers
    ws.row_dimensions[2].height = 15
    for label, sc, ec, hx in [
        ("Date",         1, 1,  GREY),
        ("Prod/Merch",   2, 3,  "D6E4F0"),
        ("Swap Dealers", 4, 5,  "D5E8D4"),
        ("Mgd Money",    6, 7,  "FFF2CC"),
        ("Other Rpt",    8, 9,  "F8D7DA"),
        ("Open Int",    10, 10, GREY),
    ]:
        if sc != ec:
            ws.merge_cells(start_row=2, start_column=sc, end_row=2, end_column=ec)
        c = ws.cell(2, sc, label)
        c.font = xfont(bold=True, color=WHITE, size=8); c.fill = xfill(NAVY)
        c.alignment = xaln(); c.border = xbdr()

    # Row 3: sub-headers
    ws.row_dimensions[3].height = 14
    for i, (label, _) in enumerate(COLS, 1):
        c = ws.cell(3, i, label)
        c.font = xfont(bold=True, color=WHITE, size=8); c.fill = xfill(BLUE)
        c.alignment = xaln(); c.border = xbdr()

    current_row = 4

    for target, data in matched.items():
        cname = SHORT_NAMES.get(target, target)

        # Contract title row
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=NCOLS)
        c = ws.cell(current_row, 1, f"— {cname} —")
        c.font = xfont(bold=True, size=10, color=WHITE)
        c.fill = xfill(BLUE); c.alignment = xaln("left"); c.border = xbdr()
        ws.row_dimensions[current_row].height = 15
        current_row += 1

        print(f"    Fetching 13-week history for {cname}...")
        hist = fetch_historical_cot(target)

        if not hist:
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row, end_column=NCOLS)
            c = ws.cell(current_row, 1, "No data available")
            c.font = xfont(italic=True, color="888888"); c.alignment = xaln("left")
            current_row += 1
        else:
            for row_i, row in enumerate(hist):
                er  = current_row + row_i
                alt = ALTROW if row_i % 2 == 0 else WHITE
                ws.row_dimensions[er].height = 13
                vals = [row["date"], row["pm_long"], row["pm_short"],
                        row["sd_long"], row["sd_short"], row["mm_long"],
                        row["mm_short"], row["or_long"], row["or_short"], row["oi"]]
                for col_i, val in enumerate(vals, 1):
                    c = ws.cell(er, col_i, val)
                    c.fill = xfill(alt); c.border = xbdr()
                    c.alignment = xaln("center" if col_i == 1 else "right")
                    if col_i == 1:
                        c.font = xfont(bold=True, size=8)
                        c.number_format = "YYYY-MM-DD"
                    else:
                        c.font = xfont(size=8)
                        c.number_format = "#,##0"
            current_row += len(hist)

        current_row += 1  # spacer

    safe_date = (report_date or "unknown").replace(" ", "_").replace(",", "")
    filename  = f"COT_{bot_label.replace(' ','_')}_{safe_date}.xlsx"
    wb.save(filename)
    print(f"    Saved: {filename}")
    return filename


def send_excel_telegram(filename, report_date, bot_label):
    """Send Excel file as Telegram document."""
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    with open(filename, "rb") as f:
        r = requests.post(url, data={
            "chat_id": TELEGRAM_CHAT_ID,
            "caption": f"📎 COT {bot_label} Data — {report_date}\nLast 13 weeks, all contracts.",
        }, files={"document": f}, timeout=30)
    if r.status_code != 200:
        print(f"Telegram file error {r.status_code}: {r.text}")
    r.raise_for_status()
    print("    Excel sent.")

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("Fetching CFTC report...")
    html = fetch_html(CFTC_URL)

    print("Parsing report...")
    parsed, report_date = parse_report(html)
    print(f"Report date: {report_date} | Contracts found: {len(parsed)}")

    if SKIP_IF_ALREADY_SAVED and already_saved(report_date):
        print(f"Report {report_date} already has saved commentary - nothing new to send.")
        return

    matched = match_contracts(parsed)
    print(f"Matched {len(matched)}/{len(TARGET_CONTRACTS)} contracts: {list(matched.keys())}")

    if not matched:
        send_telegram("⚠️ COT Bot: No target contracts matched. Check contract names.")
        return

    # Message 1: Table
    print("Sending table...")
    send_telegram(build_table_message(matched, report_date or "unknown"))
    print("Table sent.")

    # Message 2: Commentary
    print("Fetching news...")
    news = fetch_oil_news()
    print("Generating commentary...")
    commentary = build_commentary(matched, report_date or "unknown", news)
    save_commentary(matched, report_date, news, commentary)   # before Telegram, so a send error can't lose it
    send_telegram(f"🧠 *COT Commentary — {report_date}*\n\n{commentary}")
    print("Commentary sent.")

    # Message 3: Excel file
    print("Building Excel file...")
    filename = build_excel(matched, report_date or "unknown", "Petroleum")
    send_excel_telegram(filename, report_date or "unknown", "Petroleum")
    print("All done.")

if __name__ == "__main__":
    main()
